# scripts/generate_docx_manual.py
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_manual_docx():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("SERCOTEL HOTEL GUADIANA • MANUAL OPERATIVO RECEPCIÓN\n")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x16, 0x65, 0x34) # Emerald

    run_title = title_p.add_run("Gestión de Garantías y Pagos PayGold (Eurocaja Rural)")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Slate 900

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(10)
    r_intro = p_intro.add_run(
        "Procedimiento oficial para registrar reservas en MesaChef Matrix v10.3.6, generar enlaces bancarios seguros "
        "en Eurocaja Rural (Redsys PayGold), configurar la URL de notificación online automática, enviar confirmaciones con cláusula de cancelación por WhatsApp, "
        "la sincronización automática en tiempo real y la conciliación contable con Administración."
    )
    r_intro.font.size = Pt(10)
    r_intro.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # Version & Novedades Box
    nov_box = doc.add_table(rows=1, cols=1)
    nov_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    nov_cell = nov_box.rows[0].cells[0]
    nov_cell.width = Inches(6.8)
    set_cell_background(nov_cell, "F0F9FF") # Sky 50
    set_cell_margins(nov_cell, top=140, bottom=140, left=180, right=180)
    np = nov_cell.paragraphs[0]
    np.paragraph_format.space_after = Pt(4)
    nr_ver = np.add_run("✨ NOVEDADES RECIENTES DE ESTA VERSIÓN (MANUAL v10.3.6 • 28/10/2026):\n")
    nr_ver.font.bold = True
    nr_ver.font.size = Pt(10)
    nr_ver.font.color.rgb = RGBColor(0x03, 0x69, 0xA1) # Sky 700

    nov_items = [
        ("⚡ Notificación Automática Online (Webhook):", " Eurocaja Redsys actualiza la reserva a 🟢 Garantizado en vivo sin esperar al email."),
        ("📍 Captura Oficial del TPV (Paso 5):", " Se muestra la pantalla exacta de Eurocaja donde se introduce la URL del Webhook."),
        ("🔄 Botón 'Nuevo Cód.' para Modificaciones:", " Genera sufijos (GE-XXXXXX-2) si cambian personas para evitar el Error 913 (Pedido repetido)."),
        ("💰 Copia Inteligente del Importe Pendiente:", " Si ya había un pago previo, copia solo la diferencia restante."),
        ("📋 Copia Directa de URL en el Modal:", " Botón '📋 Copiar URL Webhook' en la propia ficha para agilizar el proceso."),
        ("👥 Cómputo por Personas y Cero Pagos en Lista de Puerta:", " La 'Lista de Participantes' no muestra datos monetarios ni de abono; computa personas confirmadas y personas pendientes (en vez de número de reservas), mesa y estado (✅ Confirmado / ⏳ Pendiente). Los importes y saldos quedan reservados para el 'Resumen Financiero'."),
        ("🔓 Guardado Flexible de Reserva:", " Permite guardar la reserva en blanco y añadir el enlace PayGold más tarde.")
    ]
    for bold_txt, norm_txt in nov_items:
        p_item = nov_cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(2)
        r_b = p_item.add_run("• " + bold_txt)
        r_b.font.bold = True
        r_b.font.size = Pt(9)
        r_b.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        r_n = p_item.add_run(norm_txt)
        r_n.font.size = Pt(9)
        r_n.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Table of Key Data
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Datos Clave de Acceso y Configuración")
    r_h1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    table = doc.add_table(rows=7, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    data = [
        ("Establecimiento", "Sercotel Hotel Guadiana"),
        ("Entidad Bancaria", "Eurocaja Rural (Plataforma Redsys)"),
        ("Usuario TPV Virtual", "168355840-001 (Imprescindible incluir el terminal -001)"),
        ("Portal de Acceso TPV", "https://canales.redsys.es/portal"),
        ("URL Notificación Online (Webhook)", "https://europe-west1-mesa-chef-prod.cloudfunctions.net/redsysWebhook"),
        ("Sincronización Automática", "Instantánea: MesaChef pasa a 🟢 Garantizado y anota el pago al momento"),
        ("Email Oficial de Cobro", "administracion@hotelguadiana.es (Bandeja de Isabel - Conciliación contable)")
    ]

    for idx, (label, val) in enumerate(data):
        row = table.rows[idx]
        cell_l, cell_r = row.cells[0], row.cells[1]
        cell_l.width = Inches(2.2)
        cell_r.width = Inches(4.6)
        
        set_cell_background(cell_l, "F8FAFC")
        set_cell_background(cell_r, "FFFFFF" if idx % 2 == 0 else "F8FAFC")
        set_cell_margins(cell_l)
        set_cell_margins(cell_r)

        p_l = cell_l.paragraphs[0]
        p_l.paragraph_format.space_after = Pt(2)
        r_l = p_l.add_run(label)
        r_l.font.bold = True
        r_l.font.size = Pt(9.5)
        r_l.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        p_r = cell_r.paragraphs[0]
        p_r.paragraph_format.space_after = Pt(2)
        r_r = p_r.add_run(val)
        r_r.font.size = Pt(9.5)
        if "168355840-001" in val or "https://" in val:
            r_r.font.bold = True
            r_r.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Security Callout
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = callout.rows[0].cells[0]
    c_cell.width = Inches(6.8)
    set_cell_background(c_cell, "FEF2F2") # Red 50
    set_cell_margins(c_cell, top=140, bottom=140, left=180, right=180)
    cp = c_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(2)
    cr_title = cp.add_run("🛡️ NORMATIVA DE SEGURIDAD PCI-DSS OBLIGATORIA:\n")
    cr_title.font.bold = True
    cr_title.font.size = Pt(10)
    cr_title.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)

    cr_body = cp.add_run(
        "Queda terminantemente prohibido teclear números de tarjeta completos de 16 dígitos en el recuadro "
        "de 'Observaciones'. Si el cliente facilita su tarjeta por teléfono, anota únicamente los últimos 4 dígitos "
        "(ej. 2525) en su casilla correspondiente. El pago oficial se realiza siempre mediante el enlace PayGold de Eurocaja Rural."
    )
    cr_body.font.size = Pt(9.5)
    cr_body.font.color.rgb = RGBColor(0x7F, 0x1D, 0x1D)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Step by Step
    h_steps = doc.add_heading(level=1)
    h_steps.paragraph_format.space_before = Pt(12)
    h_steps.paragraph_format.space_after = Pt(6)
    r_hsteps = h_steps.add_run("2. Procedimiento Paso a Paso")
    r_hsteps.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    steps = [
        (
            "PASO 1: Apertura de Ficha en MesaChef",
            [
                "Accede al módulo Grandes Eventos y pulsa '+ Añadir Participante' o busca la reserva existente.",
                "Rellena: Nombre Titular, Teléfono (9 dígitos), Comensales (Adultos y Niños).",
                "Fecha Cargo Previsto: Selecciona la fecha pactada límite de pago (ej. 28/12/2026).",
                "MesaChef genera automáticamente un Código de Pedido único (ej. GE-232078).",
                "Puedes guardar la ficha en cualquier momento sin necesidad de tener el enlace todavía."
            ],
            "Img/manual/step1_mesachef.png",
            "Figura 1: Ficha del participante en MesaChef con bloque seguro PayGold"
        ),
        (
            "PASO 2: Acceso al TPV Virtual de Eurocaja Rural",
            [
                "En la ficha de MesaChef, pulsa el enlace 'Abrir TPV Eurocaja Rural' (https://canales.redsys.es/portal).",
                "Usuario: 168355840-001 (muy importante incluir el terminal -001).",
                "Contraseña: Tu clave oficial del TPV de recepción y pulsa Aceptar."
            ],
            "Img/manual/step2_redsys_login.png",
            "Figura 2: Pantalla oficial de autenticación en Redsys Canales (Eurocaja Rural)"
        ),
        (
            "PASO 3: Generación del Enlace PayGold en Eurocaja",
            [
                "En el menú lateral izquierdo de Eurocaja, haz clic en PayGold (icono verde con tarjeta).",
                "Código de pedido: Pulsa 'Copiar' en MesaChef y pégalo aquí (Ctrl+V) -> Ej: GE-232078.",
                "Importe: Pulsa 'Copiar Importe' en MesaChef y pégalo aquí (Ctrl+V) -> Ej: 120.00.",
                "Descripción: Pulsa 'Copiar Descripción' en MesaChef y pégalo aquí -> Ej: Cena Fin de año 2026 en Sercotel Guadiana - Juan Pérez.",
                "Caducidad del enlace (OBLIGATORIO): Marca 'Hora exacta' e introduce la misma fecha límite del evento (ej. 28-12-2026 23:59). NUNCA dejar las 24 horas por defecto.",
                "Pulsa Aceptar / Enviar y copia el enlace seguro generado (https://sis.redsys.es/sis/p2f?t=...)."
            ],
            "Img/manual/step3_paygold_form.png",
            "Figura 3: Formulario PayGold en Eurocaja Rural con importe y fecha límite exacta"
        ),
        (
            "PASO 4: Envío del WhatsApp de Confirmación al Cliente",
            [
                "Vuelve a la ficha de MesaChef y pega el enlace en: 'Enlace PayGold generado en Eurocaja'.",
                "Haz clic en el botón verde: 'WhatsApp PayGold'.",
                "Se abre WhatsApp Web con el mensaje redactado automáticamente con el enlace seguro y la cláusula legal de no reembolso si no se abona antes de la fecha límite.",
                "Envía el mensaje al cliente y pulsa 'Guardar Ficha'. El estado pasará a 'Enlace Enviado'."
            ],
            "Img/manual/step4_whatsapp_message.png",
            "Figura 4: Mensaje oficial de WhatsApp con enlace PayGold y condiciones de cancelación"
        ),
        (
            "PASO 5: Dónde se Configura el Webhook en Eurocaja (Notificación Online)",
            [
                "Para que Eurocaja Rural notifique en tiempo real a MesaChef cada cobro, se configura una sola vez:",
                "Accede a la administración del comercio 168355840-001 en Redsys Canales.",
                "En la sección Parámetros del TPV Virtual > Notificación online:",
                "Marca la casilla: ☑️ Notificar resultado a la siguiente URL.",
                "En 'URL destino notificación', pega la URL: https://europe-west1-mesa-chef-prod.cloudfunctions.net/redsysWebhook",
                "(En MesaChef dispones del botón '📋 Copiar URL Webhook' para copiarla con 1 clic).",
                "Pulsa Aceptar / Guardar."
            ],
            "Img/manual/step_webhook_config.png",
            "Figura 5: Casilla y campo de URL destino notificación online en el portal de Eurocaja Rural"
        ),
        (
            "PASO 6: Sincronización Automática (Webhook) y Conciliación con Isabel",
            [
                "⚡ Sincronización Automática en Tiempo Real: En cuanto el cliente paga desde su móvil (Bizum o Tarjeta), Redsys notifica al Webhook de MesaChef.",
                "MesaChef actualiza la ficha solo: Cambia a '✅ Garantizado / Pagado', anota el pago con su código de autorización en el historial y guarda los últimos 4 dígitos.",
                "Recepción no necesita hacer nada a mano: la reserva se pone en verde al segundo.",
                "📑 Conciliación Contable (Isabel): Paralelamente el justificante oficial llega al email administracion@hotelguadiana.es (Isabel) para la cuadratura contable y extractos mensuales.",
                "Recepción ya no tiene que esperar a que Isabel reenvíe el correo para considerar la reserva formalizada."
            ],
            "Img/manual/step5_email_isabel.png",
            "Figura 6: Justificante oficial de cobro recibido por Isabel en administración"
        )
    ]

    for title, items, img_path, caption in steps:
        h_st = doc.add_heading(level=2)
        h_st.paragraph_format.space_before = Pt(10)
        h_st.paragraph_format.space_after = Pt(4)
        r_st = h_st.add_run(title)
        r_st.font.size = Pt(12)
        r_st.font.color.rgb = RGBColor(0x16, 0x65, 0x34)

        for it in items:
            p_it = doc.add_paragraph(style='List Bullet')
            p_it.paragraph_format.space_before = Pt(1)
            p_it.paragraph_format.space_after = Pt(2)
            r_it = p_it.add_run(it)
            r_it.font.size = Pt(9.5)
            r_it.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_img = p_img.add_run()
            run_img.add_picture(img_path, width=Inches(5.5))

            p_cap = doc.add_paragraph()
            p_cap.paragraph_format.space_after = Pt(10)
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cap = p_cap.add_run(caption)
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Modifications Section
    h_mod = doc.add_heading(level=1)
    h_mod.paragraph_format.space_before = Pt(14)
    h_mod.paragraph_format.space_after = Pt(6)
    r_hmod = h_mod.add_run("3. ¿Qué Hacer si el Cliente Modifica Comensales o Importes?")
    r_hmod.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    mod_bullets = [
        "En Redsys NUNCA se puede repetir un número de pedido ya utilizado (Error 913: Pedido repetido).",
        "Ajusta los comensales en MesaChef: si eran 2 personas (180 €) y pasan a ser 3 (270 €), el sistema calcula automáticamente el nuevo Pendiente: 90,00 €. Los pagos previos (180 €) se conservan intactos.",
        "Genera una nueva versión de pedido con 1 clic: Pulsa el botón '🔄 Nuevo Cód.' junto a Cód. Pedido. El código GE-675270 pasará automáticamente a GE-675270-2.",
        "Copia la diferencia: Pulsa '💰 Copiar Importe' (copiará solo los 90,00 € pendientes).",
        "Crea el nuevo enlace en Eurocaja: Genera el enlace PayGold con el nuevo código correlativo (GE-675270-2) y los 90 €.",
        "Envía el WhatsApp y guarda: Al pagar el cliente, el Webhook suma los nuevos 90 € al historial, completando los 270 € y dejando la reserva al 100% garantizada."
    ]
    for b in mod_bullets:
        pb = doc.add_paragraph(style='List Bullet')
        pb.paragraph_format.space_before = Pt(1)
        pb.paragraph_format.space_after = Pt(2)
        rb = pb.add_run(b)
        rb.font.size = Pt(9.5)
        rb.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Footer note
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(16)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot.add_run("MesaChef Matrix v10.3.6 • Sercotel Hotel Guadiana • Dpto. de Recepción & Administración")
    r_foot.font.size = Pt(8.5)
    r_foot.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    output_path = "MANUAL_RECEPCION_PAYGOLD_GUADIANA.docx"
    doc.save(output_path)
    print(f"Word document saved successfully to {output_path}")

if __name__ == "__main__":
    create_manual_docx()
